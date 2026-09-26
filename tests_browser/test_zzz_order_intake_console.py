"""Operational XLSX intake through the real UI, not the /demo/orders route.

All data/receipts/providers are synthetic. Tests run last to preserve earlier
source-ordered scenario assertions. No business state is inserted directly.
"""
from hashlib import sha256
from io import BytesIO
import json
from uuid import uuid4
from openpyxl import load_workbook
from playwright.sync_api import expect
from conftest import REPORTS, docker_fixture
from test_supplier_console import api, ok, eventually, eligible, suppliers, open_batch, send_ui, acknowledge_ui, tracking_file, import_ui


def intake(page):
    page.get_by_role('navigation', name='운영 메뉴').get_by_role('button', name='주문 가져오기', exact=True).click()
    expect(page.get_by_role('heading', name='주문 가져오기', exact=True)).to_be_visible()


def test_operational_file_intake_to_supplier_shipment(pages, fixture_data):
    page = pages()
    listing = next(row for row in ok(api(page, '/v1/catalog')) if row['desired_state'] == 'ACTIVE')
    intake(page)
    with page.expect_download() as download:
        page.get_by_role('link', name='빈 주문 양식', exact=True).click()
    path = REPORTS / 'synthetic-order-template.xlsx'; download.value.save_as(path)
    book = load_workbook(path); sheet = book['주문']
    assert sheet.max_row == 1
    external = '000-E2E-IMPORT-' + uuid4().hex
    sheet.append([external, '001', listing['id'], 1, listing['price'], 0,
        '합성 파일 고객', '010-0000-0000', '01234', '가상시 테스트로 123', '101호'])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'; cell.number_format = '@'
    output = BytesIO(); book.save(output); book.close(); data = output.getvalue()
    page.get_by_label('주문 마켓', exact=True).select_option(listing['marketplace'])
    page.get_by_label('주문 XLSX 파일', exact=True).set_input_files({
        'name':'synthetic-orders.xlsx', 'mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer':data})
    with page.expect_response(lambda r: r.url.endswith('/order-imports/preview') and r.request.method == 'POST') as response:
        page.get_by_role('button', name='파일 미리 보기', exact=True).click()
    preview = response.value.json(); assert response.value.status == 200 and preview['new_count'] == 1
    text = json.dumps(preview, ensure_ascii=False)
    assert all(x not in text for x in ('합성 파일 고객', '010-0000-0000', '가상시', 'pii_ciphertext'))
    submit = page.get_by_role('button', name='주문 등록 · 아직 발주 안 함', exact=True)
    expect(submit).to_be_disabled()
    source_reference = 'e2e-order-file-' + uuid4().hex
    page.get_by_label('주문 원본 참조 번호', exact=True).fill(source_reference)
    page.get_by_label('권한 있는 판매자 원본이며 미리보기의 마켓·주문·상품·금액을 확인했습니다.', exact=True).check()
    with page.expect_response(lambda r: r.url.endswith('/order-imports/commit') and r.request.method == 'POST') as response:
        submit.click()
    assert response.value.status == 201
    receipt = response.value.json(); order_id = receipt['items'][0]['id']
    before = docker_fixture('snapshot')['orders'][order_id]
    assert before['state'] == 'RECEIVED' and before['supplier_state'] is None
    assert before['payment_id'] is None and before['reservation_status'] is None
    # Re-uploading the exact same bytes is safe even with a newly generated preview ticket.
    page.get_by_role('button', name='파일 미리 보기', exact=True).click()
    expect(page.get_by_role('heading', name='미리보기 · 신규 0 / 중복 1 / 오류 0', exact=True)).to_be_visible()
    page.get_by_label('권한 있는 판매자 원본이며 미리보기의 마켓·주문·상품·금액을 확인했습니다.', exact=True).check()
    page.get_by_role('button', name='주문 등록 · 아직 발주 안 함', exact=True).click()
    expect(page.get_by_text('이미 처리한 파일입니다. 주문은 중복 생성되지 않았습니다.', exact=True)).to_be_visible()
    row = page.locator('tr').filter(has_text=external)
    row.get_by_role('button', name='원본 확인 후 검증', exact=True).click()
    page.get_by_label('현재 주문 확인 참조', exact=True).fill('e2e-current-order-' + uuid4().hex)
    page.get_by_label('현재 주문 증빙 SHA-256', exact=True).fill(sha256(b'synthetic current seller observation').hexdigest())
    verify = page.get_by_role('button', name='확인 기록 후 내부 검증', exact=True); expect(verify).to_be_disabled()
    page.get_by_label('지금 이 주문 행이 유효하고 취소되지 않았음을 원본에서 직접 확인했습니다.', exact=True).check()
    with page.expect_response(lambda r: r.url.endswith('/verify-and-validate') and r.request.method == 'POST') as response:
        verify.click()
    assert response.value.status == 200 and response.value.json()['state'] == 'SUPPLIER_ORDER_PENDING'
    command = response.value.request.post_data_json
    assert ok(api(page, f'/v1/order-intake/{order_id}/verify-and-validate', 'POST', command))['supplier_accepted'] is False
    so = next(r for r in eligible(page, [order_id]) if r['order_id'] == order_id)
    suppliers(page); page.get_by_text('새 발주 배치 생성', exact=True).click()
    page.get_by_role('combobox', name='공급사', exact=True).select_option(fixture_data['supplier_id'])
    page.get_by_role('combobox', name='고정할 프로필', exact=True).select_option(fixture_data['profile_id'])
    page.locator('label.check').filter(has_text=so['id']).get_by_role('checkbox').check()
    with page.expect_response(lambda r: r.url.endswith('/v1/supplier-order-batches') and r.request.method == 'POST') as response:
        page.get_by_role('button', name='1건 배치 생성', exact=True).click()
    assert response.value.status == 201; value = response.value.json()
    with page.expect_download() as download:
        page.get_by_role('button', name='개인정보 포함 발주서 다운로드', exact=True).click()
    download.value.save_as(REPORTS / 'synthetic-imported-supplier.xlsx')
    send_ui(page, value); state = acknowledge_ui(page, value)
    payment_id = state['items'][0]['payment_id']
    assert state['items'][0]['payment_status'] == 'EVIDENCE_PENDING'
    page.get_by_role('button', name='지급 증빙', exact=True).click()
    page.get_by_label('은행·공급사 지급 참조 번호', exact=True).fill('e2e-import-payment-' + uuid4().hex)
    page.get_by_label('증빙 SHA-256', exact=True).fill(sha256(b'synthetic imported order payment').hexdigest())
    page.get_by_role('button', name='증빙 기록 · 아직 지급 확정 아님', exact=True).click()
    expect(page.get_by_text('관리자 확인 대기 중입니다. 운영자는 지급 확정을 수행할 수 없습니다.', exact=True)).to_be_visible()
    admin = pages('admin'); open_batch(admin, value['id'])
    admin.get_by_role('button', name='지급 증빙', exact=True).click()
    admin.get_by_label('해당 수취처로 이 금액의 실제 지급이 완료되었음을 증빙으로 확인했습니다.', exact=True).check()
    admin.get_by_role('button', name='관리자 지급 증빙 확인', exact=True).click()
    expect(admin.get_by_text('관리자 지급 증빙 확인이 완료되었습니다. 이 화면에서 송금을 실행한 것은 아닙니다.', exact=True)).to_be_visible()
    tracking = tracking_file(so['id'], external, '007777777777')
    import_ui(page, fixture_data['profile_id'], tracking)
    result = eventually(lambda: docker_fixture('snapshot')['orders'][order_id], lambda x: x['shipment_id'] is not None and x['shipment_job_statuses']==['DEAD'])
    assert not result['marketplace_synced']
    work = ok(api(admin, '/v1/operations'))
    shipment = next(x for x in work['shipments'] if x['order_id']==order_id)
    confirmation = {k:shipment[k] for k in ('marketplace','external_order_id','external_line_id','tracking_hash')}
    confirmation.update(reference='synthetic-intake-shipment-confirmation',evidence_hash=sha256(b'synthetic-intake-shipment').hexdigest(),confirmed=True)
    first = ok(api(admin,f"/v1/shipments/{shipment['id']}/external-confirmation",'POST',confirmation))
    assert first == ok(api(admin,f"/v1/shipments/{shipment['id']}/external-confirmation",'POST',confirmation))
    result = docker_fixture('snapshot')['orders'][order_id]
    assert result['manual_shipment_confirmation']
    assert result['payment_id'] == payment_id and result['payment_status'] == 'SUCCEEDED'
    assert result['supplier_state'] == 'SHIPPED' and result['shipment_jobs'] == 1
    (REPORTS / 'order-intake-result.json').write_text(json.dumps({
        'synthetic_only':True, 'order_id':order_id, 'import_id':receipt['id'], 'file_sha256':sha256(data).hexdigest(),
        'before_verification':before, 'completed':result, 'real_money_moved':False}, indent=2))


def test_intake_viewer_has_only_safe_history(pages):
    page = pages('viewer'); intake(page)
    expect(page.get_by_label('주문 XLSX 파일', exact=True)).not_to_be_visible()
    expect(page.get_by_role('button', name='원본 확인 후 검증', exact=True)).not_to_be_visible()
    history = ok(api(page, '/v1/order-imports')); assert history
    assert all('items' not in row and 'source_reference' not in row for row in history)
    assert api(page, '/v1/order-intake/pending')['status'] == 403
    assert api(page, '/v1/order-intake/unknown/verify-and-validate', 'POST', {})['status'] == 403
    assert api(page, '/v1/order-imports/template')['status'] == 403
