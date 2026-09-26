"""Browser acceptance against actual built containers; only providers are DEMO.

UI actions are used for the complete normal transaction. API setup in the
focused negative cases still uses real HTTP, authentication and application
services; no route.fulfill, mocks, or direct business-state mutation is used.
"""
from hashlib import sha256
from io import BytesIO
import json
import time
from uuid import uuid4
import pytest
from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect
from conftest import REPORTS, ORIGIN, docker_fixture


def api(page, path, method='GET', payload=None, *, csrf=True):
    return page.evaluate('''async ({path,method,payload,csrf}) => {
      const headers = {'Content-Type':'application/json'};
      if(csrf) headers['X-CSRF-Token'] = decodeURIComponent(document.cookie.split('; ').find(x=>x.startsWith('saup_csrf='))?.split('=')[1] || '');
      const r=await fetch('/api'+path,{method,headers,credentials:'same-origin',cache:'no-store',...(payload===null?{}:{body:JSON.stringify(payload)})});
      return {status:r.status,body:await r.json()};
    }''', {'path':path, 'method':method, 'payload':payload, 'csrf':csrf})


def ok(response):
    assert 200 <= response['status'] < 300, response
    return response['body']


def eventually(fn, predicate, seconds=20):
    until = time.monotonic() + seconds
    while True:
        value = fn()
        if predicate(value): return value
        if time.monotonic() >= until: raise AssertionError('Expected worker/business state was not reached')
        time.sleep(0.6)


def batch(page, batch_id):
    return ok(api(page, f'/v1/supplier-order-batches/{batch_id}'))


def suppliers(page):
    page.get_by_role('navigation', name='운영 메뉴').get_by_role('button', name='공급사 운영', exact=True).click()
    expect(page.get_by_role('heading', name='공급사 Excel 운영', exact=True)).to_be_visible()


def open_batch(page, batch_id):
    suppliers(page)
    page.get_by_role('button', name=batch_id, exact=True).click()
    expect(page.get_by_role('heading', name=batch_id, exact=True)).to_be_visible()


def eligible(page, order_ids):
    return eventually(lambda: ok(api(page, '/v1/supplier-orders/eligible')),
        lambda rows: len([r for r in rows if r['order_id'] in order_ids]) == len(order_ids))


def create_by_api(page, fixture_data, path='MANUAL_EVIDENCE'):
    item = next(row for row in ok(api(page, '/v1/catalog')) if row['desired_state'] == 'ACTIVE')
    raw = {'marketplace': item['marketplace'], 'external_id': 'E2E-'+uuid4().hex, 'external_line_id': '1',
           'listing_id': item['id'], 'quantity': 1, 'gross_sale': item['price'],
           'address': {'recipient': '합성 고객', 'phone': '010-0000-0000', 'postal_code': '01234',
                       'address1': '가상시 테스트로 123', 'address2': '101호'}}
    order = ok(api(page, '/demo/orders', 'POST', raw))
    so = next(r for r in eligible(page, [order['id']]) if r['order_id'] == order['id'])
    value = ok(api(page, '/v1/supplier-order-batches', 'POST', {
        'supplier_id': fixture_data['supplier_id'], 'profile_id': fixture_data['profile_id'],
        'supplier_order_ids': [so['id']], 'payment_path': path, 'idempotency_key': 'e2e-'+uuid4().hex}))
    return value, order, so


def send_ui(page, value):
    page.get_by_label('전송 참조 번호', exact=True).fill('e2e-send-'+value['id'])
    button = page.get_by_role('button', name='전송 사실 기록', exact=True)
    expect(button).to_be_disabled()
    page.get_by_label('이 배치의 파일을 실제로 전송했습니다.', exact=True).check()
    button.click()
    expect(page.get_by_role('heading', name='공급사 접수 결과 기록', exact=True)).to_be_visible()
    assert batch(page, value['id'])['status'] == 'SENT'


def acknowledge_ui(page, value, *, reject_id=None, changed_amount=None):
    for line in value['items']:
        group = page.get_by_role('group', name=line['supplier_order_id'], exact=True)
        group.get_by_role('combobox', name='실제 공급사 응답', exact=True).select_option('ORDER_NOT_ACCEPTED' if line['supplier_order_id'] == reject_id else 'ACCEPTED')
        if changed_amount is not None:
            group.get_by_text('공급사가 변경 조건을 통보했습니다', exact=True).click()
            group.get_by_label('통보 총액', exact=True).fill(str(changed_amount))
    page.get_by_label('공급사 확인 참조 번호', exact=True).fill('e2e-ack-'+value['id'])
    page.get_by_label('실제 공급사 응답을 확인했으며 위 결과를 그대로 기록합니다.', exact=True).check()
    page.get_by_role('button', name='접수 결과 기록', exact=True).click()
    expect(page.get_by_role('heading', name='공급사 접수 결과 기록', exact=True)).not_to_be_visible()
    return batch(page, value['id'])


def tracking_file(so_id, external_id, tracking):
    book = Workbook(); sheet = book.active; sheet.title = 'Sheet1'
    sheet.append(['발주ID', '마켓주문ID', '택배사', '송장번호'])
    sheet.append([so_id, external_id, 'CJ', tracking])
    for row in sheet:
        for cell in row: cell.data_type = 's'; cell.number_format = '@'
    result = BytesIO(); book.save(result); book.close(); return result.getvalue()


def import_ui(page, profile_id, content):
    page.get_by_role('navigation', name='운영 메뉴').get_by_role('button', name='파일', exact=True).click()
    # The Files panel has one profile selector.
    page.get_by_role('combobox', name='공급사 엑셀 프로필', exact=True).select_option(profile_id)
    with page.expect_response(lambda r: f'/v1/imports/{profile_id}/shipment' in r.url and r.request.method == 'POST') as response:
        page.get_by_label('송장 가져오기', exact=True).set_input_files({'name':'synthetic-tracking.xlsx',
            'mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':content})
    assert response.value.status == 202
    return response.value.json()


def test_manual_evidence_supplier_transaction(pages, fixture_data):
    page = pages()
    created = []
    for _ in range(2):
        with page.expect_response(lambda r: r.url.endswith('/demo/orders') and r.request.method == 'POST') as response:
            page.get_by_role('button', name='＋ 모의 주문', exact=True).click()
        assert response.value.status == 202
        created.append(response.value.json()['id'])
    rows = [r for r in eligible(page, created) if r['order_id'] in created]
    suppliers(page); page.get_by_text('새 발주 배치 생성', exact=True).click()
    page.get_by_role('combobox', name='공급사', exact=True).select_option(fixture_data['supplier_id'])
    page.get_by_role('combobox', name='고정할 프로필', exact=True).select_option(fixture_data['profile_id'])
    for row in rows: page.locator('label.check').filter(has_text=row['id']).get_by_role('checkbox').check()
    with page.expect_response(lambda r: r.url.endswith('/v1/supplier-order-batches') and r.request.method == 'POST') as response:
        page.get_by_role('button', name='2건 배치 생성', exact=True).click()
    assert response.value.status == 201
    value = response.value.json(); create_command = response.value.request.post_data_json
    expect(page.get_by_role('heading', name=value['id'], exact=True)).to_be_visible()
    assert value['status'] == 'FILE_READY' and value['profile_version'] == fixture_data['profile_version']
    assert value['sent_at'] is None and value['acknowledged_at'] is None
    assert ok(api(page, '/v1/supplier-order-batches', 'POST', create_command))['id'] == value['id']
    copies = []
    for index in range(2):
        with page.expect_download() as download:
            page.get_by_role('button', name='개인정보 포함 발주서 다운로드', exact=True).click()
        target = REPORTS / f'synthetic-supplier-{index}.xlsx'; download.value.save_as(target)
        copies.append(target.read_bytes())
    assert copies[0] == copies[1] and sha256(copies[0]).hexdigest() == value['file_hash']
    book = load_workbook(BytesIO(copies[0]), data_only=False)
    cells = list(book.active); headers = [c.value for c in cells[0]]
    assert len(cells) == 3
    assert {row[0].value for row in cells[1:]} == {r['id'] for r in rows}
    for row in cells[1:]:
        assert row[headers.index('우편번호')].value == '01234'
        assert row[headers.index('전화번호')].value == '010-0000-0000'
        assert row[headers.index('공급사SKU')].data_type == 's'
        assert all(c.data_type != 'f' for c in row)
    book.close()
    state = batch(page, value['id'])
    assert state['status'] == 'EXPORTED' and state['sent_at'] is None
    send_ui(page, value)
    send = {'send_channel':'EMAIL', 'send_reference':'e2e-send-'+value['id'], 'file_hash':value['file_hash']}
    assert ok(api(page, f"/v1/supplier-order-batches/{value['id']}/mark-sent", 'POST', send))['status'] == 'SENT'
    conflict = api(page, f"/v1/supplier-order-batches/{value['id']}/mark-sent", 'POST', {**send, 'send_reference':'e2e-conflict'})
    assert conflict['status'] == 409 and conflict['body']['error'] == 'SUPPLIER_SEND_CONFLICT'
    rejected_id = value['items'][1]['supplier_order_id']
    state = acknowledge_ui(page, value, reject_id=rejected_id)
    assert state['accepted_count'] == 1 and state['rejected_count'] == 1
    accepted = next(x for x in state['items'] if x['ack_status'] == 'ACCEPTED')
    rejected = next(x for x in state['items'] if x['ack_status'] == 'REJECTED')
    assert accepted['payment_status'] == 'EVIDENCE_PENDING' and rejected['payment_id'] is None
    ack = {'accepted_order_ids':[accepted['supplier_order_id']], 'rejected':[{'supplier_order_id':rejected_id,'reason':'ORDER_NOT_ACCEPTED'}], 'reference':'e2e-ack-'+value['id']}
    ok(api(page, f"/v1/supplier-order-batches/{value['id']}/acknowledge", 'POST', ack))
    page.get_by_role('button', name='지급 증빙', exact=True).click()
    page.get_by_label('은행·공급사 지급 참조 번호', exact=True).fill('e2e-payment-'+value['id'])
    local_proof = b'SYNTHETIC ONLY - not a bank receipt or actual money movement'
    page.get_by_label('로컬 증빙 파일에서 해시 계산 · 파일은 업로드하지 않음', exact=True).set_input_files({
        'name':'synthetic-evidence.txt','mimeType':'text/plain','buffer':local_proof})
    expect(page.get_by_label('증빙 SHA-256', exact=True)).to_have_value(sha256(local_proof).hexdigest())
    page.get_by_role('button', name='증빙 기록 · 아직 지급 확정 아님', exact=True).click()
    expect(page.get_by_text('관리자 확인 대기 중입니다. 운영자는 지급 확정을 수행할 수 없습니다.', exact=True)).to_be_visible()
    payment = ok(api(page, f"/v1/supplier-payments/{accepted['payment_id']}"))
    assert payment['status'] == 'EVIDENCE_PENDING' and payment['evidence_status'] != 'CONFIRMED'
    confirmation = {'amount':payment['amount'],'destination_fingerprint':payment['destination_fingerprint'],
                    'reference':payment['evidence']['reference'],'confirmed_money_moved':True}
    forbidden = api(page, f"/v1/supplier-payment-evidence/{payment['evidence_id']}/confirm", 'POST', confirmation)
    assert forbidden['status'] == 403 and forbidden['body']['error'] == 'ROLE_FORBIDDEN'
    admin = pages('admin'); open_batch(admin, value['id'])
    admin.get_by_role('button', name='지급 증빙', exact=True).click()
    confirm_button = admin.get_by_role('button', name='관리자 지급 증빙 확인', exact=True)
    expect(confirm_button).to_be_disabled()
    admin.get_by_label('해당 수취처로 이 금액의 실제 지급이 완료되었음을 증빙으로 확인했습니다.', exact=True).check()
    confirm_button.click()
    expect(confirm_button).not_to_be_visible()
    assert ok(api(admin, f"/v1/supplier-payments/{accepted['payment_id']}"))['status'] == 'SUCCEEDED'
    ok(api(admin, f"/v1/supplier-payment-evidence/{payment['evidence_id']}/confirm", 'POST', confirmation))
    page.reload(); open_batch(page, value['id'])
    page.get_by_role('button', name='지급 증빙', exact=True).click()
    expect(page.get_by_text('관리자 지급 증빙 확인이 완료되었습니다. 이 화면에서 송금을 실행한 것은 아닙니다.', exact=True)).to_be_visible()
    expect(page.get_by_text('관리자 확인 대기 중입니다. 운영자는 지급 확정을 수행할 수 없습니다.', exact=True)).not_to_be_visible()
    external = next(o for o in ok(api(page, '/v1/orders')) if o['id'] == accepted['order_id'])['external_id']
    content = tracking_file(accepted['supplier_order_id'], external, '001234567890')
    first = import_ui(page, fixture_data['profile_id'], content)
    assert import_ui(page, fixture_data['profile_id'], content) == first
    snapshot = eventually(lambda: docker_fixture('snapshot'), lambda x: x['orders'][accepted['order_id']]['marketplace_synced'])
    observed = snapshot['orders'][accepted['order_id']]
    assert observed['supplier_state'] == 'SHIPPED' and observed['shipment_jobs'] == 1
    assert observed['confirmation_id'] and observed['evidence_id'] and snapshot['balanced_journals']
    assert snapshot['orders'][rejected['order_id']]['payment_id'] is None
    for event in ('SUPPLIER_BATCH_CREATED','SUPPLIER_BATCH_FILE_EXPORTED','SUPPLIER_BATCH_SENT',
                  'SUPPLIER_BATCH_ACKNOWLEDGED','SUPPLIER_ORDER_ACCEPTED','SUPPLIER_ORDER_REJECTED',
                  'SUPPLIER_PAYMENT_EVIDENCE_RECORDED','SUPPLIER_PAYMENT_CONFIRMED','MARKETPLACE_SHIPMENT_UPDATED'):
        assert snapshot['audit_counts'].get(event, 0) > 0
    assert snapshot['counts']['payments'] == 1 and snapshot['counts']['supplier_payment_evidence'] == 1
    assert snapshot['counts']['supplier_payment_confirmations'] == 1 and snapshot['counts']['shipments'] == 1
    (REPORTS / 'manual-transaction.json').write_text(json.dumps(snapshot, indent=2, sort_keys=True))
    open_batch(page, value['id']); expect(page.locator('.supplier-detail')).to_contain_text('SHIPPED')


def test_viewer_status_only_and_no_pii(pages):
    page = pages('viewer'); suppliers(page)
    expect(page.get_by_text('새 발주 배치 생성', exact=True)).not_to_be_visible()
    expect(page.get_by_role('button', name='송장 파일 가져오기', exact=True)).not_to_be_visible()
    values = ok(api(page, '/v1/supplier-order-batches'))
    for value in values:
        detail = ok(api(page, f"/v1/supplier-order-batches/{value['id']}"))
        serialized = json.dumps(detail, ensure_ascii=False)
        for protected in ('pii_ciphertext','file_ciphertext','original_address','010-0000-0000','가상시 테스트로'):
            assert protected not in serialized
    identifier = values[0]['id'] if values else 'not-authorized-resource'
    file_result = page.context.request.get(ORIGIN+f'/api/v1/supplier-order-batches/{identifier}/file')
    assert file_result.status == 403
    denied = api(page, '/v1/supplier-order-batches', 'POST', {})
    assert denied['status'] == 403 and denied['body']['error'] == 'ROLE_FORBIDDEN'
    if values:
        open_batch(page, identifier)
        expect(page.get_by_role('button', name='개인정보 포함 발주서 다운로드', exact=True)).not_to_be_visible()
        expect(page.get_by_role('button', name='배치 취소 요청', exact=True)).not_to_be_visible()


def test_real_csrf_origin_and_session_cookie_contract(pages):
    page = pages()
    denied = api(page, '/v1/supplier-order-batches', 'POST', {}, csrf=False)
    assert denied['status'] == 403 and denied['body']['error'] == 'CSRF_REJECTED'
    csrf = next(c['value'] for c in page.context.cookies() if c['name'] == 'saup_csrf')
    response = page.context.request.post(ORIGIN+'/api/v1/supplier-order-batches',
        headers={'Origin':'https://untrusted.example','X-CSRF-Token':csrf}, data={})
    assert response.status == 403 and response.json()['error'] == 'ORIGIN_REJECTED'
    cookie = next(c for c in page.context.cookies() if c['name'] == 'saup_session')
    assert cookie['httpOnly'] and cookie['sameSite'] == 'Strict'
    assert 'saup_session' not in page.evaluate('document.cookie')
    response = page.context.request.get(ORIGIN+'/api/v1/supplier-order-batches')
    assert response.headers['cache-control'] == 'no-store'


def test_supplier_price_change_visible_and_no_payment(pages, fixture_data):
    page = pages(); value, order, so = create_by_api(page, fixture_data)
    open_batch(page, value['id']); send_ui(page, value)
    state = acknowledge_ui(page, value, changed_amount=so['amount']+1)
    assert state['items'][0]['status'] == 'MANUAL_REVIEW' and state['items'][0]['amount'] == so['amount']
    assert state['items'][0]['payment_id'] is None
    expect(page.locator('.supplier-detail')).to_contain_text('MANUAL_REVIEW')
    expect(page.get_by_role('button', name='지급 증빙', exact=True)).not_to_be_visible()
    snapshot = docker_fixture('snapshot')
    assert snapshot['orders'][order['id']]['payment_id'] is None


@pytest.mark.parametrize('after_send', [False, True], ids=['before-send','after-send'])
def test_cancellation_handoff(pages, fixture_data, after_send):
    page = pages(); value, order, _so = create_by_api(page, fixture_data)
    open_batch(page, value['id'])
    if after_send: send_ui(page, value)
    page.once('dialog', lambda d: d.accept())
    page.get_by_role('button', name='배치 취소 요청', exact=True).click()
    expect(page.get_by_role('button', name='개인정보 포함 발주서 다운로드', exact=True)).not_to_be_visible()
    state = batch(page, value['id'])
    snapshot = docker_fixture('snapshot')
    observed = snapshot['orders'][order['id']]
    if after_send:
        assert state['status'] == 'CANCEL_PENDING' and observed['reservation_status'] == 'HELD'
        admin = pages('admin'); open_batch(admin, value['id'])
        admin.on('dialog', lambda d: d.accept('e2e-cancel-proof') if d.type == 'prompt' else d.accept())
        admin.get_by_role('button', name='공급사 취소 확인', exact=True).click()
        expect(admin.get_by_role('button', name='공급사 취소 확인', exact=True)).not_to_be_visible()
        observed = docker_fixture('snapshot')['orders'][order['id']]
    assert observed['supplier_state'] == 'CANCELLED' and observed['reservation_status'] == 'RELEASED'
    assert observed['payment_id'] is None


def test_demo_provider_path_still_explicit_and_ships(pages, fixture_data):
    page = pages(); value, order, so = create_by_api(page, fixture_data, path='DEMO_PROVIDER')
    open_batch(page, value['id']); expect(page.locator('.supplier-detail')).to_contain_text('DEMO_PROVIDER')
    send_ui(page, value); acknowledge_ui(page, value)
    eventually(lambda: batch(page, value['id']), lambda x: x['items'][0]['status'] == 'SHIPMENT_PENDING')
    external = next(o for o in ok(api(page, '/v1/orders')) if o['id'] == order['id'])['external_id']
    content = tracking_file(so['id'], external, '001234567891')
    import_ui(page, fixture_data['profile_id'], content)
    state = eventually(lambda: docker_fixture('snapshot'), lambda x: x['orders'][order['id']]['marketplace_synced'])
    observed = state['orders'][order['id']]
    assert observed['supplier_state'] == 'SHIPPED' and observed['payment_status'] == 'SUCCEEDED'
    assert observed['evidence_id'] is None and observed['confirmation_id'] is None and observed['shipment_jobs'] == 1
    (REPORTS / 'demo-provider-transaction.json').write_text(json.dumps(state, indent=2, sort_keys=True))
