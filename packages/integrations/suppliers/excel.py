"""Bounded profile-driven XLSX adapter. No formulas, macros, remote links or coercive IDs."""
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from packages.domain.errors import DomainError
from packages.domain.schemas import PriceRow, ShipmentRow


class ExcelProfile(BaseModel):
    model_config = ConfigDict(extra="forbid")
    sheet: str = "Sheet1"
    header_row: int = Field(default=1, ge=1, le=20)
    columns: dict[str, str]
    defaults: dict = Field(default_factory=dict)
    sku_map: dict[str, str] = Field(default_factory=dict)
    date_formats: list[str] = Field(default_factory=lambda: ["%Y-%m-%d"])
    order_columns: dict[str, str] = Field(default_factory=lambda: {
        "supplier_order_id": "발주ID", "marketplace_order_id": "마켓주문ID",
        "supplier_sku": "공급사SKU", "quantity": "수량", "recipient": "수령인",
        "phone": "전화번호", "postal_code": "우편번호", "address1": "주소", "address2": "상세주소"})
    shipment_columns: dict[str, str] = Field(default_factory=lambda: {
        "supplier_order_id": "발주ID", "marketplace_order_id": "마켓주문ID",
        "courier": "택배사", "tracking": "송장번호"})

    @model_validator(mode="after")
    def mappings(self):
        for mapping in (self.columns, self.order_columns, self.shipment_columns):
            if not mapping or len(mapping) != len(set(mapping.values())):
                raise ValueError("DUPLICATE_PROFILE_HEADERS")
        if not set(PriceRow.model_fields).issubset(set(self.columns) | set(self.defaults)):
            raise ValueError("MISSING_PRICE_MAPPING")
        if set(self.shipment_columns) != set(ShipmentRow.model_fields):
            raise ValueError("MISSING_SHIPMENT_MAPPING")
        needed = {"supplier_order_id", "marketplace_order_id", "supplier_sku", "quantity", "recipient", "phone", "postal_code", "address1", "address2"}
        if set(self.order_columns) != needed:
            raise ValueError("INVALID_ORDER_MAPPING")
        return self


@dataclass
class ImportResult:
    rows: list[tuple[int, dict]] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def number(value, nullable=False):
    if value is None and nullable:
        return None
    if isinstance(value, bool) or value is None:
        raise ValueError("INVALID_NUMBER")
    try:
        # Comma grouping is accepted only as conventional thousands groups.
        if isinstance(value, str) and "," in value:
            import re
            if not re.fullmatch(r"[0-9]{1,3}(,[0-9]{3})+", value):
                raise ValueError("INVALID_NUMBER")
            value = value.replace(",", "")
        d = Decimal(str(value))
        if not d.is_finite() or d != d.to_integral_value() or not 0 <= d <= 10**12:
            raise ValueError("INVALID_NUMBER")
        return int(d)
    except InvalidOperation as exc:
        raise ValueError("INVALID_NUMBER") from exc


def validate_archive(data: bytes):
    if not data or len(data) > 5_000_000:
        raise DomainError("FILE_SIZE_LIMIT", 422)
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 150 or sum(i.file_size for i in infos) > 25_000_000:
                raise DomainError("XLSX_EXPANSION_LIMIT", 422)
            for info in infos:
                name = info.filename.lower()
                if info.flag_bits & 1 or "vbaproject" in name or "externallinks/" in name:
                    raise DomainError("UNSAFE_XLSX_CONTENT", 422)
                if info.file_size > max(1, info.compress_size) * 300:
                    raise DomainError("XLSX_EXPANSION_LIMIT", 422)
                if ".." in name.split("/"):
                    raise DomainError("UNSAFE_XLSX_CONTENT", 422)
    except BadZipFile as exc:
        raise DomainError("INVALID_XLSX", 422) from exc


def parse(data: bytes, profile: ExcelProfile, kind: str) -> ImportResult:
    validate_archive(data)
    if kind not in {"price", "shipment"}:
        raise DomainError("UNKNOWN_IMPORT_KIND", 422)
    mapping = profile.columns if kind == "price" else profile.shipment_columns
    model = PriceRow if kind == "price" else ShipmentRow
    result = ImportResult()
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise DomainError("INVALID_XLSX", 422) from exc
    try:
        if profile.sheet not in workbook.sheetnames:
            raise DomainError("MISSING_SHEET", 422)
        sheet = workbook[profile.sheet]
        if (sheet.max_row or 0) > 5020 or (sheet.max_column or 0) > 64:
            raise DomainError("WORKSHEET_LIMIT", 422)
        iterator = sheet.iter_rows(min_row=profile.header_row)
        cells = next(iterator, ())
        headers = [c.value for c in cells]
        used = [x for x in headers if x is not None]
        if any(c.data_type == "f" for c in cells) or len(used) != len(set(used)):
            raise DomainError("INVALID_HEADERS", 422)
        if not set(mapping.values()).issubset(set(headers)):
            raise DomainError("MISSING_COLUMNS", 422)
        indexes = {field: headers.index(header) for field, header in mapping.items()}
        for rownum, cells in enumerate(iterator, start=profile.header_row + 1):
            if rownum > profile.header_row + 5000:
                raise DomainError("WORKSHEET_LIMIT", 422)
            if all(c.value is None for c in cells):
                continue
            if any(c.data_type == "f" for c in cells):
                result.errors.append({"row": rownum, "code": "FORMULA_NOT_ALLOWED"})
                continue
            raw = dict(profile.defaults) if kind == "price" else {}
            raw.update({key: cells[index].value if index < len(cells) else None for key, index in indexes.items()})
            try:
                idfields = ("supplier_sku",) if kind == "price" else ("supplier_order_id", "marketplace_order_id", "tracking")
                if any(not isinstance(raw.get(k), str) or not raw[k].strip() for k in idfields):
                    raise ValueError("IDENTIFIER_MUST_BE_TEXT")
                if kind == "price":
                    for key in ("cost", "shipping", "stock", "weight_grams"):
                        raw[key] = number(raw.get(key), key == "stock")
                validated = model.model_validate(raw)
                result.rows.append((rownum, validated.model_dump()))
            except ValidationError:
                result.errors.append({"row": rownum, "code": "ROW_VALIDATION_FAILED"})
            except ValueError as exc:
                result.errors.append({"row": rownum, "code": str(exc)})
        keyfield = "supplier_sku" if kind == "price" else "supplier_order_id"
        counts = Counter(row[keyfield] for _, row in result.rows)
        duplicate_rows = {i for i, row in result.rows if counts[row[keyfield]] > 1}
        if kind == "shipment":
            tracking_counts = Counter((row["courier"], row["tracking"]) for _, row in result.rows)
            duplicate_rows |= {i for i, row in result.rows if tracking_counts[row["courier"], row["tracking"]] > 1}
        result.errors.extend({"row": i, "code": "DUPLICATE_ROW"} for i in sorted(duplicate_rows))
        result.rows = [(i, row) for i, row in result.rows if i not in duplicate_rows]
        return result
    finally:
        workbook.close()


def workbook_bytes(headers: list[str], rows: list[list], sheet_name="Sheet1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for values in [headers, *rows]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Preserve exact text without interpreting =,+,-,@ as formulas.
                cell.number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163A46")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(38, max(14, max(len(str(c.value or "")) for c in column) + 3))
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.properties.creator = "Saup"
    workbook.properties.created = datetime(2000, 1, 1)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def export_orders(profile: ExcelProfile, orders: list[dict]) -> bytes:
    ids = [row["supplier_order_id"] for row in orders]
    if len(ids) != len(set(ids)):
        raise DomainError("DUPLICATE_EXPORT_ORDER", 422)
    for row in orders:
        from packages.domain.schemas import Address
        Address.model_validate({k: row[k] for k in Address.model_fields})
        if number(row["quantity"]) == 0:
            raise DomainError("INVALID_QUANTITY", 422)
    return workbook_bytes(list(profile.order_columns.values()),
        [[row[key] for key in profile.order_columns] for row in orders], profile.sheet)
