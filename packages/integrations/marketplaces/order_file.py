"""Saup's versioned canonical order template, NOT a guessed marketplace format.

Reuses the existing bounded XLSX archive and text-safe writer. No formulas,
macros, external links, type-coerced IDs or silently ignored populated columns.
"""
from collections import Counter
from io import BytesIO
from zipfile import ZipFile
import re
from defusedxml.ElementTree import iterparse, fromstring
from openpyxl import load_workbook
from pydantic import ValidationError
from packages.domain.errors import DomainError
from packages.domain.schemas import Address, OrderInput
from packages.integrations.suppliers.excel import ImportResult, number, validate_archive, workbook_bytes

VERSION = "SAUP_ORDER_XLSX_V1"
SHEET = "주문"
MAX_ROWS = 500
COLUMNS = {
    "external_id": "마켓주문번호", "external_line_id": "주문행번호", "listing_id": "판매상품ID",
    "quantity": "수량", "gross_sale": "할인전주문금액", "discount": "할인금액",
    "recipient": "수령인", "phone": "전화번호", "postal_code": "우편번호",
    "address1": "주소", "address2": "상세주소",
}
ID_FIELDS = {"external_id", "external_line_id", "listing_id", "phone", "postal_code"}


def template():
    # Headers only: no synthetic customer or fake product becomes a real order.
    return workbook_bytes(list(COLUMNS.values()), [], SHEET)


def _check_physical_cells(data):
    """Do not trust XLSX dimension metadata to hide out-of-bounds rows/columns."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.namelist()) != len(set(archive.namelist())):
                raise DomainError("UNSAFE_XLSX_CONTENT", 422)
            sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml")]
            if len(sheets) != 1:
                raise DomainError("ORDER_TEMPLATE_SHEET_REQUIRED", 422)
            # Only the one scanned worksheet may be loaded. A nonstandard
            # relationship must not redirect the parser to unscanned XML.
            relationships = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = []
            for rel in relationships:
                if rel.attrib.get("Type", "").endswith("/worksheet"):
                    target = rel.attrib.get("Target", "")
                    target = target.lstrip("/") if target.startswith("/") else "xl/" + target
                    if rel.attrib.get("TargetMode") == "External" or target != sheets[0]:
                        raise DomainError("UNSAFE_XLSX_CONTENT", 422)
                    targets.append(target)
            if targets != sheets:
                raise DomainError("ORDER_TEMPLATE_SHEET_REQUIRED", 422)
            with archive.open(sheets[0]) as xml:
                count, current_row, last_row, last_column = 0, None, 0, ""
                for event, element in iterparse(xml, events=("start", "end")):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if event == "start" and tag == "row":
                        index = element.attrib.get("r", "")
                        if current_row is not None or not index.isdigit() or not last_row < int(index) <= MAX_ROWS+1:
                            raise DomainError("ORDER_WORKSHEET_LIMIT", 422)
                        current_row, last_column = int(index), ""
                    elif event == "start" and tag == "c":
                        count += 1
                        match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", element.attrib.get("r", ""))
                        if (not match or current_row is None or len(match[1]) != 1 or
                                not last_column < match[1] <= "K" or int(match[2]) != current_row or
                                count > (MAX_ROWS+1)*len(COLUMNS)):
                            raise DomainError("ORDER_WORKSHEET_LIMIT", 422)
                        last_column = match[1]
                    elif event == "end":
                        if tag == "row":
                            last_row, current_row = current_row, None
                        element.clear()
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("INVALID_XLSX", 422) from exc


def parse_orders(data: bytes, marketplace: str) -> ImportResult:
    validate_archive(data)
    _check_physical_cells(data)
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise DomainError("INVALID_XLSX", 422) from exc
    result = ImportResult()
    try:
        if wb.sheetnames != [SHEET]:
            raise DomainError("ORDER_TEMPLATE_SHEET_REQUIRED", 422)
        ws = wb[SHEET]
        if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > len(COLUMNS):
            raise DomainError("ORDER_WORKSHEET_LIMIT", 422)
        # Explicit iteration bounds also cap documents with absent/false dimensions.
        iterator = ws.iter_rows(min_row=1, max_row=MAX_ROWS + 2, max_col=len(COLUMNS) + 1)
        headers = next(iterator, ())
        names = [x.value for x in headers[:len(COLUMNS)]]
        if (any(x.data_type == "f" for x in headers) or names != list(COLUMNS.values()) or
                any(x.value is not None for x in headers[len(COLUMNS):])):
            raise DomainError("ORDER_TEMPLATE_HEADERS_REQUIRED", 422)
        for row_number, cells in enumerate(iterator, 2):
            if all(x.value is None for x in cells):
                continue
            if row_number > MAX_ROWS + 1 or cells[-1].value is not None:
                raise DomainError("ORDER_WORKSHEET_LIMIT", 422)
            if any(x.data_type == "f" for x in cells):
                result.errors.append({"row": row_number, "code": "FORMULA_NOT_ALLOWED"})
                continue
            raw = dict(zip(COLUMNS, (x.value for x in cells)))
            try:
                for key in ID_FIELDS:
                    value = raw[key]
                    if not isinstance(value, str) or not value.strip():
                        raise ValueError("IDENTIFIER_MUST_BE_TEXT")
                    if any(ord(c) < 32 or ord(c) == 127 for c in value):
                        raise ValueError("INVALID_IDENTIFIER")
                address = {k: raw[k] for k in Address.model_fields}
                if address["address2"] is None:
                    address["address2"] = ""  # Absence, not address rewriting.
                address = Address.model_validate(address).model_dump()
                order = OrderInput.model_validate({"marketplace": marketplace,
                    **{k: raw[k] for k in ("external_id", "external_line_id", "listing_id")},
                    **{k: number(raw[k]) for k in ("quantity", "gross_sale", "discount")}, "address": address})
                result.rows.append((row_number, order.model_dump()))
            except ValidationError as exc:
                # Never include Pydantic's 'input' values or exception string (PII).
                fields = sorted({str(e["loc"][0]) for e in exc.errors(include_input=False) if e["loc"]})
                result.errors.append({"row": row_number, "code": "ROW_VALIDATION_FAILED", "fields": fields})
            except ValueError as exc:
                code = str(exc)
                if code not in {"IDENTIFIER_MUST_BE_TEXT", "INVALID_IDENTIFIER", "INVALID_NUMBER"}:
                    code = "ROW_VALIDATION_FAILED"
                result.errors.append({"row": row_number, "code": code})
        if not result.rows and not result.errors:
            raise DomainError("EMPTY_ORDER_FILE", 422)
        counts = Counter((raw["external_id"], raw["external_line_id"]) for _, raw in result.rows)
        duplicate_rows = {i for i, raw in result.rows if counts[(raw["external_id"], raw["external_line_id"])] > 1}
        result.errors.extend({"row": i, "code": "DUPLICATE_ORDER_ROW"} for i in sorted(duplicate_rows))
        result.rows = [(i, raw) for i, raw in result.rows if i not in duplicate_rows]
        return result
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("INVALID_XLSX", 422) from exc
    finally:
        wb.close()
