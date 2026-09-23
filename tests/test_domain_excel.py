from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from packages.domain.errors import DomainError
from packages.domain.pricing import required_price, margin, krw
from packages.domain.state_machine import validate_transition, NORMAL
from packages.domain.schemas import Address
from packages.domain.discovery import evaluate, experiment_state
from packages.infrastructure.security import hash_password, verify_password, Cipher
from packages.integrations.suppliers.excel import ExcelProfile, parse, workbook_bytes, export_orders
from cryptography.fernet import Fernet


def profile():
    return ExcelProfile(columns={"supplier_sku":"SKU", "title":"상품명", "cost":"원가", "shipping":"배송비", "stock":"재고"},
        defaults={"category":"농산물", "origin":"대한민국", "tax_type":"EXEMPT", "weight_grams":3000, "grade":"공급사 표기", "unit":"상자"})


def xlsx(rows):
    return workbook_bytes(["SKU", "상품명", "원가", "배송비", "재고"], rows)

@pytest.mark.parametrize("cost,fee,target", [(13300,'.1','.15'),(1,'.333333','.1'),(0,'0','0'),(99999,'.2','.3')])
def test_price_clears_rounded_margin(cost,fee,target):
    price = required_price(cost,Decimal(fee),Decimal(target))
    assert price % 100 == 0
    assert margin(price,cost,Decimal(fee)).percentage >= Decimal(target)

@pytest.mark.parametrize("bad", [True,-1,1.5,Decimal(5),'5',10**13])
def test_money_is_bounded_integer(bad):
    with pytest.raises(DomainError): krw(bad)

def test_impossible_margin():
    with pytest.raises(DomainError): required_price(10,Decimal('.6'),Decimal('.5'))

@pytest.mark.parametrize("before,after",list(zip(NORMAL,NORMAL[1:])))
def test_normal_state_transitions(before,after):
    validate_transition(before,after)

@pytest.mark.parametrize("before,after", [('RECEIVED','SHIPPED'),('CLOSED','RECEIVED'),('CANCELLED','SUPPLIER_PAID')])
def test_illegal_transitions(before,after):
    with pytest.raises(DomainError): validate_transition(before,after)

def test_original_address_preserved():
    a={'recipient':' 가상 고객 ', 'phone':'010-0000-0000','postal_code':'01234','address1':'가상시 테스트로 123','address2':' 101호 '}
    assert Address.model_validate(a).model_dump() == a

def test_password_and_encryption():
    hashed=hash_password('correct-test-password')
    assert verify_password('correct-test-password',hashed)
    assert not verify_password('wrong',hashed)
    cipher=Cipher(Fernet.generate_key().decode())
    text='원본 주소 (공백 보존)'
    sealed=cipher.encrypt({'address':text})
    assert text not in sealed
    assert cipher.decrypt(sealed)['address'] == text

def test_valid_price_profile_and_text_sku():
    result=parse(xlsx([['00123','가상 감귤','13,300',3000,10]]),profile(),'price')
    assert result.errors == []
    assert result.rows[0][1]['supplier_sku']=='00123'
    assert result.rows[0][1]['cost']==13300

@pytest.mark.parametrize('sku,cost,stock', [(None,1,1),(123,1,1),('001',-5,1),('001',1.5,1),('001',1,-1),('001','NaN',1),('001','1,23',1)])
def test_bad_price_rows_are_isolated(sku,cost,stock):
    result=parse(xlsx([[sku,'bad',cost,1,stock],['GOOD','ok',100,1,10]]),profile(),'price')
    assert len(result.errors)==1
    assert len(result.rows)==1

def test_all_duplicates_rejected():
    result=parse(xlsx([['A','a',1,1,1],['A','b',2,1,1]]),profile(),'price')
    assert result.rows==[] and len(result.errors)==2

def test_formula_not_evaluated():
    data=xlsx([['A','a',1,1,1]])
    book=load_workbook(BytesIO(data));book.active['C2']='=1+1'
    out=BytesIO();book.save(out)
    result=parse(out.getvalue(),profile(),'price')
    assert result.errors[0]['code']=='FORMULA_NOT_ALLOWED'

def test_malformed_file():
    with pytest.raises(DomainError): parse(b'not-an-excel-file',profile(),'price')

def test_missing_headers():
    with pytest.raises(DomainError): parse(workbook_bytes(['A'],[['a']]),profile(),'price')

def test_export_text_not_formula():
    row={'supplier_order_id':'s1','marketplace_order_id':'m1','supplier_sku':'001','quantity':1,
      'recipient':'=Customer', 'phone':'010-0000-0000','postal_code':'01234','address1':'가상시 테스트로 123','address2':'@room'}
    data=export_orders(profile(),[row])
    book=load_workbook(BytesIO(data),data_only=False)
    assert book.active['E2'].value=='=Customer'
    assert book.active['E2'].data_type=='s'
    assert book.active['G2'].value=='01234'

def test_duplicate_export_rejected():
    with pytest.raises(DomainError): export_orders(profile(),[{'supplier_order_id':'x'},{'supplier_order_id':'x'}])

def test_discovery_not_low_reviews_only():
    result=evaluate({'purchases':0,'review_velocity':0,'search_interest':0,'seasonality':50},
      {'review_barrier':0,'seller_concentration':0,'price_pressure':0,'sameness':0},Decimal('.2'),'manual export',50)
    assert result.action=='WATCH'

def test_experiment_never_scales_only_revenue():
    assert experiment_state(5000,40,-1000,0)=='REWORK'
    assert experiment_state(5000,40,10000,0)=='SCALE'
