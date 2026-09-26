"""Phase 10 behavioural tests: synthetic data only, no real money/provider calls."""
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import json
import pytest
from sqlalchemy import select, func
from openpyxl import load_workbook
from packages.domain.errors import DomainError
from packages.domain.supplier_operations import Actor
from packages.infrastructure.models import (
    Supplier, SupplierOrder, SupplierOrderIntent, SupplierOrderBatch, SupplierOrderBatchItem,
    SupplierExcelProfile, SupplierProduct, MarketplaceListing, Order, Payment, Reservation,
    SupplierBatchAcknowledgement, SupplierPaymentEvidence, SupplierPaymentConfirmation,
    SupplierCancellation, OperationalIntervention, Review, AuditEvent, Job, Shipment, Posting, Journal,
)
from packages.application.files import Files
from packages.application.finance import post, balance
from packages.application.common import lock_treasury
from packages.integrations.suppliers.excel import ExcelProfile, workbook_bytes

OP = Actor("operator", "operator")
ADMIN = Actor("admin", "admin")
VIEWER = Actor("reader", "viewer")

@pytest.fixture
def excel_env(env):
    c = env["commerce"]
    c.clock = lambda: datetime(2026, 9, 25, 3, 0, tzinfo=timezone.utc)
    with env["factory"].begin() as s:
        s.get(Supplier, env["ids"]["supplier_id"]).mode = "excel"
        for row in s.scalars(select(SupplierProduct)):
            row.last_inventory_at = c.clock()
    env["ops"] = c.supplier_operations
    return env


def make_order(env, raw, *, name="supplier-001", index=0):
    data = deepcopy(raw); data["external_id"] = name
    data["listing_id"] = env["ids"]["listing_ids"][index]
    with env["factory"]() as s:
        data["gross_sale"] = s.get(MarketplaceListing, data["listing_id"]).price
    result = env["commerce"].ingest(data)
    env["commerce"].process_order(result["id"])
    with env["factory"]() as s:
        so = s.scalar(select(SupplierOrder).where(SupplierOrder.order_id == result["id"]))
        assert so is not None
        return so.id


def batch(env, ids, *, path="MANUAL_EVIDENCE", key="batch-001", **changes):
    payload = {"supplier_id": env["ids"]["supplier_id"], "profile_id": env["ids"]["profile_id"],
        "supplier_order_ids": ids, "idempotency_key": key, "payment_path": path, **changes}
    return env["ops"].create_batch(payload, OP)


def send(env, b, **changes):
    return env["ops"].mark_sent(b["id"], {"send_channel": "EMAIL", "send_reference": "supplier-send-001",
        "file_hash": b["file_hash"], **changes}, OP)


def ack(env, b, *, accepted=None, rejected=None, changes=None, reference="supplier-ack-001"):
    if accepted is None:
        accepted = [x["supplier_order_id"] for x in b["items"]]
    return env["ops"].acknowledge(b["id"], {"accepted_order_ids": accepted, "rejected": rejected or [],
        "reported_changes": changes or [], "reference": reference}, OP)


def payment(env, so_id):
    with env["factory"]() as s:
        return s.scalar(select(Payment).where(Payment.supplier_order_id == so_id))


def accepted_order(env, raw, *, path="MANUAL_EVIDENCE"):
    so_id = make_order(env, raw)
    b = batch(env, [so_id], path=path)
    send(env, b); ack(env, b)
    return so_id, b, payment(env, so_id)


def proof_payload(env, p):
    data = env["ops"].payment_details(p.id, OP)
    method = "SUPPLIER_DEPOSIT" if data["deposit_amount"] == data["amount"] else "MANUAL_TRANSFER" if not data["deposit_amount"] else "OTHER_APPROVED_METHOD"
    return {key: data[key] for key in ("supplier_id", "amount", "bank_amount", "deposit_amount", "destination_fingerprint")} | {
        "method": method, "reference": "bank-receipt-001", "evidence_hash": "a" * 64}


def record(env, p, **changes):
    return env["ops"].record_evidence(p.id, proof_payload(env, p) | changes, OP)


def confirm(env, p, evidence, *, actor=ADMIN, **changes):
    return env["ops"].confirm_evidence(evidence["id"], {"amount": p.amount,
        "destination_fingerprint": p.destination_fingerprint, "reference": "bank-receipt-001",
        "confirmed_money_moved": True, **changes}, actor)


def tracking(env, so_id, number="001234567890"):
    with env["factory"]() as s:
        so = s.get(SupplierOrder, so_id); order = s.get(Order, so.order_id)
        return {"supplier_order_id": so_id, "marketplace_order_id": order.external_id,
            "courier": "CJ", "tracking": number}


def test_pending_intent_does_not_claim_file_exists(excel_env, order_input):
    e = excel_env; so_id = make_order(e, order_input)
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so_id).status == "PENDING"
        intent = s.scalar(select(SupplierOrderIntent))
        assert intent and intent.amount == s.get(SupplierOrder, so_id).amount
        assert not s.scalar(select(Payment))
        assert not s.scalar(select(SupplierOrderBatch))
        assert "phone" not in json.dumps(intent.snapshot)


def test_batch_membership_idempotent_payload_bound_and_nonempty(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input)
    a = batch(e, [so]); b = batch(e, [so])
    assert a["id"] == b["id"] and b["order_count"] == 1
    with pytest.raises(DomainError):
        batch(e, [so], key="incompatible-batch")
    with pytest.raises(DomainError, match="IDEMPOTENCY_CONFLICT"):
        batch(e, [so], path="DEMO_PROVIDER")
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        batch(e, [], key="empty")
    with pytest.raises(DomainError, match="DUPLICATE_BATCH_ORDER"):
        batch(e, [so, so], key="duplicate")
    with e["factory"]() as s:
        assert s.scalar(select(func.count()).select_from(SupplierOrderBatchItem)) == 1


def test_cross_supplier_and_profile_rejected(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input)
    with e["factory"].begin() as s:
        other = Supplier(name="other", mode="excel", cutoff="23:59", destination_approved=True, destination_fingerprint="b"*64)
        s.add(other); s.flush()
        profile = s.get(SupplierExcelProfile, e["ids"]["profile_id"])
        p2 = SupplierExcelProfile(supplier_id=other.id, name="other-profile", version=2, mapping=profile.mapping)
        s.add(p2); s.flush(); sid, pid = other.id, p2.id
    with pytest.raises(DomainError, match="CROSS_SUPPLIER_BATCH"):
        batch(e, [so], supplier_id=sid, profile_id=pid)
    with pytest.raises(DomainError, match="BATCH_SUPPLIER_PROFILE_MISMATCH"):
        batch(e, [so], profile_id=pid)


@pytest.mark.parametrize("cancel_at", ["before_batch", "after_generation", "after_export"])
def test_cancel_before_send_invalidates_old_file_and_releases(excel_env, order_input, cancel_at):
    e = excel_env; so = make_order(e, order_input); b = None
    if cancel_at != "before_batch":
        b = batch(e, [so])
    if cancel_at == "after_export":
        e["ops"].export_batch(b["id"], OP)
    first = e["ops"].cancel_order(so, OP)
    assert first == e["ops"].cancel_order(so, OP)
    with e["factory"]() as s:
        row = s.get(SupplierOrder, so)
        assert row.status == "CANCELLED"
        assert s.get(Order, row.order_id).state == "CANCELLED"
        assert s.scalar(select(Reservation)).status == "RELEASED"
    with pytest.raises(DomainError):
        batch(e, [so], key="cancelled-order")
    if b:
        with pytest.raises(DomainError, match="BATCH_FILE_INVALIDATED"):
            e["ops"].export_batch(b["id"], OP)
        with pytest.raises(DomainError):
            send(e, b)


def test_unsent_cancel_retains_historical_membership_allows_survivor_rebatch(excel_env, order_input):
    e = excel_env
    one = make_order(e, order_input, name="one"); two = make_order(e, order_input, name="two")
    b = batch(e, [one, two]); e["ops"].export_batch(b["id"], OP)
    e["ops"].cancel_order(one, OP)
    b2 = batch(e, [two], key="replacement")
    assert b2["order_count"] == 1
    with e["factory"]() as s:
        assert s.scalar(select(func.count()).select_from(SupplierOrderBatchItem)) == 3
        assert s.get(SupplierOrderBatch, b["id"]).status == "INVALIDATED"


def test_export_frozen_profile_exact_bytes_and_text_safety(excel_env, order_input):
    e = excel_env; raw = deepcopy(order_input)
    raw["address"]["recipient"] = "=DANGEROUS()"
    raw["address"]["address1"] = "  가상시 테스트로 123  "
    one = make_order(e, raw, name="000000123"); make_order(e, order_input, name="excluded")
    b = batch(e, [one])
    with e["factory"].begin() as s:
        p = s.get(SupplierExcelProfile, e["ids"]["profile_id"])
        p.version = 8; p.mapping = dict(p.mapping) | {"sheet": "Changed"}
    data, name = e["ops"].export_batch(b["id"], OP)
    repeat, _ = e["ops"].export_batch(b["id"], OP)
    assert repeat == data and hashlib.sha256(data).hexdigest() == b["file_hash"]
    book = load_workbook(BytesIO(data)); sheet = book.active
    assert sheet.max_row == 2 and sheet.title != "Changed"
    values = [c.value for c in sheet[2]]
    assert "000000123" in values and "01234" in values and "=DANGEROUS()" in values
    assert "  가상시 테스트로 123  " in values
    assert all(c.data_type != "f" for row in sheet for c in row)
    book.close()
    with e["factory"]() as s:
        saved = s.get(SupplierOrderBatch, b["id"])
        assert saved.profile_version == 1
        assert "DANGEROUS" not in saved.file_ciphertext
        assert s.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.event == "SUPPLIER_BATCH_FILE_EXPORTED")) == 2
        assert not s.scalar(select(Payment))
    assert e["ops"].get_batch(b["id"])["status"] == "EXPORTED"


def test_mark_sent_idempotent_conflict_and_no_acceptance(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so])
    send(e, b); send(e, b)
    with pytest.raises(DomainError, match="SUPPLIER_SEND_CONFLICT"):
        send(e, b, send_reference="conflicting-send")
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so).status == "SENT"
        assert not s.scalar(select(Payment))
        assert s.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.event == "SUPPLIER_BATCH_SENT")) == 1


def test_ack_full_acceptance_and_duplicate_creates_one_payment(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input)
    assert p.status == "EVIDENCE_PENDING"
    ack(e, b); send(e, b)
    with e["factory"]() as s:
        assert s.scalar(select(func.count()).select_from(Payment)) == 1
        assert s.scalar(select(func.count()).select_from(SupplierBatchAcknowledgement)) == 1
        assert s.get(SupplierOrder, so).status == "PAYMENT_PENDING"
        assert not s.scalar(select(Job).where(Job.kind == "payment.execute"))
        assert s.scalar(select(Review).where(Review.category == "SUPPLIER_FILE_ACK_REQUIRED")).status == "RESOLVED"
    with pytest.raises(DomainError, match="SUPPLIER_ACK_CONFLICT"):
        ack(e, b, reference="conflict")


@pytest.mark.parametrize("case", ["unknown", "cross_batch", "missing", "duplicate", "accepted_and_rejected"])
def test_ack_membership_rejected(excel_env, order_input, case):
    e = excel_env; one = make_order(e, order_input, name="one"); two = make_order(e, order_input, name="two")
    b = batch(e, [one]); batch(e, [two], key="other-batch"); send(e, b)
    accepted = {"unknown": ["unknown"], "cross_batch": [two], "missing": [], "duplicate": [one, one], "accepted_and_rejected": [one]}[case]
    rejected = [{"supplier_order_id": one, "reason": "OTHER"}] if case == "accepted_and_rejected" else []
    with pytest.raises(DomainError, match="ACK_BATCH_MEMBERSHIP_MISMATCH"):
        ack(e, b, accepted=accepted, rejected=rejected)
    assert payment(e, one) is None


def test_ack_requires_real_send(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); e["ops"].export_batch(b["id"], OP)
    with pytest.raises(DomainError, match="SUPPLIER_BATCH_NOT_SENT"):
        ack(e, b)


@pytest.mark.parametrize("full", [False, True])
def test_partial_and_full_rejection_blocks_payment(excel_env, order_input, full):
    e = excel_env; one = make_order(e, order_input, name="one"); two = make_order(e, order_input, name="two")
    b = batch(e, [one, two]); send(e, b)
    rejected = [one, two] if full else [two]
    accepted = [] if full else [one]
    result = ack(e, b, accepted=accepted, rejected=[{"supplier_order_id": x, "reason": "ORDER_NOT_ACCEPTED"} for x in rejected])
    assert result["rejected_count"] == len(rejected)
    for so in rejected:
        assert payment(e, so) is None
        with e["factory"]() as s:
            assert s.get(SupplierOrder, so).status == "REJECTED"
    if not full:
        assert payment(e, one).status == "EVIDENCE_PENDING"


@pytest.mark.parametrize("change", [
    {"amount": 19000}, {"shipping": 9000}, {"quantity": 2}, {"sku": "CHANGED"}, {"destination_changed": True},
])
def test_changed_supplier_terms_never_mutate_frozen_amount_or_pay(excel_env, order_input, change):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); send(e, b)
    with e["factory"]() as s:
        old_amount = s.get(SupplierOrder, so).amount
    result = ack(e, b, changes=[{"supplier_order_id": so, **change}])
    assert result["review_count"] == 1 and payment(e, so) is None
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so).amount == old_amount
        assert s.scalar(select(Review).where(Review.category == "SUPPLIER_TERMS_CHANGED"))
        assert s.scalar(select(Reservation)).status == "HELD"


def test_price_changed_reason_without_new_amount_blocks_and_retains_reservation(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); send(e, b)
    ack(e, b, accepted=[], rejected=[{"supplier_order_id": so, "reason": "PRICE_CHANGED"}])
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so).status == "MANUAL_REVIEW"
        assert s.scalar(select(Reservation)).status == "HELD"
    assert payment(e, so) is None


@pytest.mark.parametrize("condition", ["destination", "inactive", "cost", "cash", "cancel"])
def test_acceptance_payment_preparation_rechecks_current_safety(excel_env, order_input, condition):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); send(e, b)
    with e["factory"].begin() as s:
        supplier = s.get(Supplier, e["ids"]["supplier_id"])
        if condition == "destination": supplier.destination_fingerprint = "c" * 64
        if condition == "inactive": supplier.active = False
        if condition == "cost": s.scalar(select(SupplierProduct).where(SupplierProduct.supplier_id == supplier.id)).cost += 1
        if condition == "cash":
            lock_treasury(s); post(s, "cash-withdrawal", "TEST", {"BANK": -900000, "EQUITY": 900000})
    if condition == "cancel": e["ops"].cancel_order(so, OP)
    ack(e, b)
    assert payment(e, so) is None


def test_evidence_record_is_not_payment_success_and_confirmation_is_idempotent(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input)
    evidence = record(e, p); assert evidence == record(e, p)
    assert payment(e, so).status == "EVIDENCE_PENDING"
    result = confirm(e, p, evidence); assert result == confirm(e, p, evidence)
    with e["factory"]() as s:
        assert s.get(Payment, p.id).status == "SUCCEEDED"
        assert s.get(SupplierOrder, so).status == "SHIPMENT_PENDING"
        assert s.scalar(select(Reservation)).status == "SPENT"
        assert s.scalar(select(func.count()).select_from(SupplierPaymentEvidence)) == 1
        assert s.scalar(select(func.count()).select_from(SupplierPaymentConfirmation)) == 1
        events = list(s.scalars(select(AuditEvent).where(AuditEvent.event == "SUPPLIER_PAYMENT_CREATED")))
        assert events[0].new_value["source"] == "MANUAL_EVIDENCE" and events[0].new_value["demo"] is False
        assert not s.scalar(select(Job).where(Job.kind == "payment.execute"))


def test_evidence_mismatch_review_persists_after_failure(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input)
    with pytest.raises(DomainError, match="PAYMENT_EVIDENCE_SNAPSHOT_MISMATCH"):
        record(e, p, amount=p.amount+1, bank_amount=p.amount+1)
    with e["factory"]() as s:
        assert s.scalar(select(Review).where(Review.category == "PAYMENT_EVIDENCE_SNAPSHOT_MISMATCH"))
        assert not s.scalar(select(SupplierPaymentEvidence))
    assert payment(e, so).status == "EVIDENCE_PENDING"


def test_confirmation_requires_admin_and_matching_metadata(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input); evidence = record(e, p)
    with pytest.raises(DomainError, match="FORBIDDEN"):
        confirm(e, p, evidence, actor=OP)
    with pytest.raises(DomainError, match="PAYMENT_CONFIRMATION_SNAPSHOT_MISMATCH"):
        confirm(e, p, evidence, amount=p.amount+1)
    with e["factory"]() as s:
        assert s.scalar(select(Review).where(Review.category == "PAYMENT_CONFIRMATION_SNAPSHOT_MISMATCH"))
    assert payment(e, so).status == "EVIDENCE_PENDING"
    confirm(e, p, evidence)
    with pytest.raises(DomainError, match="PAYMENT_CONFIRMATION_CONFLICT"):
        confirm(e, p, evidence, reference="different")


def test_deposit_is_supplier_specific_and_balanced(excel_env, order_input):
    e = excel_env; sid = e["ids"]["supplier_id"]
    with e["factory"].begin() as s:
        lock_treasury(s)
        other = Supplier(name="unrelated-supplier", mode="excel")
        s.add(other); s.flush(); other_id = other.id
        post(s, "deposit-own", "TEST", {f"DEPOSIT:{sid}": 20000, "BANK": -20000})
        post(s, "deposit-other", "TEST", {f"DEPOSIT:{other_id}": 50000, "BANK": -50000})
    so, b, p = accepted_order(e, order_input)
    assert proof_payload(e, p)["method"] == "SUPPLIER_DEPOSIT"
    confirm(e, p, record(e, p))
    with e["factory"]() as s:
        assert balance(s, f"DEPOSIT:{other_id}") == 50000
        assert balance(s, f"DEPOSIT:{sid}") == 20000-p.amount >= 0
        for journal in s.scalars(select(Journal)):
            assert s.scalar(select(func.sum(Posting.delta)).where(Posting.journal_id == journal.id)) == 0


@pytest.mark.parametrize("limit", ["single", "daily"])
def test_limits_remain_gated(excel_env, order_input, limit):
    e = excel_env
    if limit == "single": e["settings"].auto_payment_limit = 1
    else: e["settings"].daily_payment_limit = 1
    so, b, p = accepted_order(e, order_input)
    assert p.status == "MANUAL_APPROVAL"
    if limit == "daily":
        with pytest.raises(DomainError, match="DAILY_PAYMENT_LIMIT"):
            e["commerce"].approve_payment(p.id, p.amount, p.destination_fingerprint, "admin")
        return
    e["commerce"].approve_payment(p.id, p.amount, p.destination_fingerprint, "admin")
    assert payment(e, so).status == "EVIDENCE_PENDING"
    confirm(e, p, record(e, p))


def test_lowered_limit_at_confirmation_blocks_with_review(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input); evidence = record(e, p)
    e["settings"].daily_payment_limit = 1
    with pytest.raises(DomainError, match="DAILY_PAYMENT_LIMIT"):
        confirm(e, p, evidence)
    assert payment(e, so).status == "EVIDENCE_PENDING"


def test_manual_payment_cannot_be_completed_by_provider_finalizer(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input)
    with pytest.raises(DomainError, match="PAYMENT_NOT_PROVIDER_CONFIRMABLE"):
        e["commerce"].finalize_payment(p.id, "fake-provider")
    e["commerce"].execute_payment(p.id)
    assert payment(e, so).status == "EVIDENCE_PENDING"


def test_demo_provider_excel_path_full_shipment_excel_import(excel_env, order_input):
    e = excel_env; so, b, p = accepted_order(e, order_input, path="DEMO_PROVIDER")
    e["worker"].drain()
    assert payment(e, so).status == "SUCCEEDED"
    row = tracking(e, so)
    with e["factory"]() as s:
        mapping = ExcelProfile.model_validate(s.get(SupplierExcelProfile, e["ids"]["profile_id"]).mapping)
    data = workbook_bytes(list(mapping.shipment_columns.values()), [[row[key] for key in mapping.shipment_columns]])
    files = Files(e["commerce"])
    first = files.submit(e["ids"]["profile_id"], "shipment", data)
    assert first == files.submit(e["ids"]["profile_id"], "shipment", data)
    e["worker"].drain()
    with e["factory"].begin() as s:
        lock_treasury(s)
        e["commerce"].add_shipment(s, row, e["ids"]["supplier_id"])
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so).status == "SHIPPED"
        assert s.scalar(select(Shipment)).marketplace_synced
        assert s.scalar(select(func.count()).select_from(Shipment)) == 1
        assert s.scalar(select(func.count()).select_from(Job).where(Job.kind == "marketplace.shipment")) == 1


@pytest.mark.parametrize("state", ["unacknowledged", "rejected", "unpaid", "cancelled"])
def test_ineligible_supplier_order_cannot_ship(excel_env, order_input, state):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); send(e, b)
    if state == "rejected": ack(e, b, accepted=[], rejected=[{"supplier_order_id": so, "reason": "OTHER"}])
    if state == "unpaid": ack(e, b)
    if state == "cancelled": e["ops"].cancel_order(so, OP)
    with e["factory"].begin() as s:
        lock_treasury(s)
        with pytest.raises(DomainError): e["commerce"].add_shipment(s, tracking(e, so))
    with e["factory"]() as s:
        assert not s.scalar(select(Shipment))


@pytest.mark.parametrize("stage", ["sent", "acknowledged", "evidence_recorded", "paid"])
def test_post_send_cancellation_retains_reservations_until_supplier_confirmation(excel_env, order_input, stage):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); send(e, b)
    if stage != "sent": ack(e, b)
    p = payment(e, so)
    if stage in {"evidence_recorded", "paid"}:
        proof = record(e, p)
        if stage == "paid": confirm(e, p, proof)
    e["ops"].cancel_order(so, OP)
    with e["factory"]() as s:
        assert s.get(SupplierOrder, so).status == "CANCEL_PENDING"
        assert s.scalar(select(Reservation)).status == ("SPENT" if stage == "paid" else "HELD")
        assert not s.scalar(select(Job).where(Job.kind == "supplier.cancel"))
    payload = {"reference": "supplier-cancel-001", "supplier_confirmed_cancelled": True}
    first = e["ops"].confirm_cancellation(so, payload, ADMIN)
    assert first == e["ops"].confirm_cancellation(so, payload, ADMIN)
    with e["factory"]() as s:
        reservation = s.scalar(select(Reservation))
        if stage in {"paid", "evidence_recorded"}:
            assert first["status"] == "FINANCIAL_REVIEW"
            assert reservation.status in {"HELD", "SPENT"}
        else:
            assert first["status"] == "CONFIRMED" and reservation.status == "RELEASED"
    with pytest.raises(DomainError, match="SUPPLIER_CANCELLATION_CONFLICT"):
        e["ops"].confirm_cancellation(so, payload | {"reference": "changed"}, ADMIN)


def test_audit_and_manual_measurement_reconstructs_full_workflow(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so]); e["ops"].export_batch(b["id"], OP)
    send(e, b); ack(e, b); p = payment(e, so); confirm(e, p, record(e, p))
    with e["factory"].begin() as s:
        lock_treasury(s); e["commerce"].add_shipment(s, tracking(e, so))
    e["worker"].drain(); e["ops"].resolve_batch(b["id"], ADMIN)
    with e["factory"]() as s:
        events = set(s.scalars(select(AuditEvent.event)))
        assert {"SUPPLIER_BATCH_CREATED", "SUPPLIER_BATCH_FILE_EXPORTED", "SUPPLIER_BATCH_SENT",
            "SUPPLIER_BATCH_ACKNOWLEDGED", "SUPPLIER_ORDER_ACCEPTED", "SUPPLIER_PAYMENT_EVIDENCE_RECORDED",
            "SUPPLIER_PAYMENT_CONFIRMED", "SUPPLIER_BATCH_RESOLVED"} <= events
        for audit in s.scalars(select(AuditEvent)):
            assert "010-0000" not in json.dumps(audit.new_value) and "가상시" not in json.dumps(audit.new_value)
    metrics = e["ops"].metrics()
    assert metrics["production_automation_percentage"] is None
    assert metrics["interventions_by_category"]["SUPPLIER_FILE_SENT_MANUALLY"] == 1


def test_viewer_service_authority(excel_env, order_input):
    e = excel_env; so = make_order(e, order_input); b = batch(e, [so])
    with pytest.raises(DomainError, match="FORBIDDEN"):
        e["ops"].export_batch(b["id"], VIEWER)
    with pytest.raises(DomainError, match="FORBIDDEN"):
        e["ops"].cancel_order(so, VIEWER)
